"""Render a resume dict to a clean, single-page, ATS-parsable .docx and convert it to PDF."""
import logging
import math
import shutil
import subprocess
import sys
from pathlib import Path
from typing import List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

logger = logging.getLogger(__name__)

PAGE_W_IN, PAGE_H_IN = 8.27, 11.69  # A4
MARGIN_IN = 0.55
FONT = "Calibri"
FONT_SIZES = [10.5, 10.0, 9.5, 9.0, 8.5, 8.0]  # largest size that still fits one page is chosen

Block = Tuple[str, str, str]  # (kind, bold_prefix, text)


def build_blocks(resume: dict) -> List[Block]:
    """Layout-independent list of resume blocks (shared by the renderer and the page-fit estimator)."""
    blocks: List[Block] = [
        ("name", "", resume.get("name", "")),
        ("title", "", resume.get("title", "")),
    ]
    if resume.get("contact"):
        blocks.append(("contact", "", resume["contact"]))

    blocks += [("heading", "", "Summary"), ("para", "", resume.get("summary", ""))]

    blocks.append(("heading", "", "Skills"))
    for group in resume.get("skills", []):
        blocks.append(("skill", group["category"] + ": ", ", ".join(group["items"])))

    blocks += [("heading", "", "Education"), ("para", "", resume.get("education", ""))]

    if resume.get("experience_bullets"):
        blocks.append(("heading", "", "Experience"))
        if resume.get("experience_heading"):
            blocks.append(("subhead", "", resume["experience_heading"]))
        blocks += [("bullet", "", b["text"]) for b in resume["experience_bullets"]]
        if resume.get("experience_tech_stack"):
            blocks.append(("meta", "Tech Stack: ", ", ".join(resume["experience_tech_stack"])))

    if resume.get("projects"):
        blocks.append(("heading", "", "Projects"))
        for project in resume["projects"]:
            blocks.append(("subhead", "", project["name"]))
            blocks += [("bullet", "", b["text"]) for b in project["bullets"]]
            if project.get("tech_stack"):
                blocks.append(("meta", "Tech Stack: ", ", ".join(project["tech_stack"])))

    if resume.get("certifications"):
        blocks.append(("heading", "", "Certifications"))
        blocks += [("bullet", "", c) for c in resume["certifications"]]
    return blocks


def _style(kind: str, size: float) -> dict:
    """Font size and paragraph spacing (points) per block kind."""
    return {
        "name": dict(size=size + 6, before=0, after=1),
        "title": dict(size=size + 0.5, before=0, after=1),
        "contact": dict(size=size - 0.5, before=0, after=2),
        "heading": dict(size=size + 0.5, before=6, after=2),
        "para": dict(size=size, before=0, after=1),
        "skill": dict(size=size, before=0, after=0.5),
        "subhead": dict(size=size, before=3, after=0.5),
        "bullet": dict(size=size, before=0, after=0.5),
        "meta": dict(size=size - 0.5, before=0, after=1),
    }[kind]


def estimate_height_pt(blocks: List[Block], size: float) -> float:
    usable_w_pt = (PAGE_W_IN - 2 * MARGIN_IN) * 72
    total = 0.0
    for kind, prefix, text in blocks:
        style = _style(kind, size)
        indent = 14 if kind == "bullet" else 0
        chars_per_line = max((usable_w_pt - indent) / (style["size"] * 0.50), 10)
        lines = max(1, math.ceil((len(prefix) + len(text)) / chars_per_line))
        total += lines * style["size"] * 1.22 + style["before"] + style["after"]
    return total


def choose_font_size(blocks: List[Block]) -> float:
    usable_h_pt = (PAGE_H_IN - 2 * MARGIN_IN) * 72 - 6
    for size in FONT_SIZES:
        if estimate_height_pt(blocks, size) <= usable_h_pt:
            return size
    return FONT_SIZES[-1]


def _bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for attr, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "444444")):
        bottom.set(qn(attr), value)
    borders.append(bottom)
    p_pr.append(borders)


def _run(paragraph, text: str, size: float, bold=False, italic=False, color: Optional[RGBColor] = None):
    run = paragraph.add_run(text)
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold, run.italic = bold, italic
    if color is not None:
        run.font.color.rgb = color
    return run


def render_docx(resume: dict, path: Path) -> Path:
    """Write the resume as a single-column .docx (no tables/text boxes, so ATS parsers read it cleanly)."""
    blocks = build_blocks(resume)
    base_size = choose_font_size(blocks)

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(PAGE_W_IN), Inches(PAGE_H_IN)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(MARGIN_IN))
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(base_size)

    for kind, prefix, text in blocks:
        style = _style(kind, base_size)
        if kind == "bullet":
            paragraph = doc.add_paragraph(style="List Bullet")
        else:
            paragraph = doc.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(style["before"]), Pt(style["after"]), 1.0
        size = style["size"]

        if kind == "name":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(paragraph, text, size, bold=True)
        elif kind in ("title", "contact"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(paragraph, text, size, italic=(kind == "title"))
        elif kind == "heading":
            _run(paragraph, text.upper(), size, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))
            _bottom_border(paragraph)
        elif kind == "subhead":
            _run(paragraph, text, size, bold=True)
        elif kind == "meta":
            _run(paragraph, prefix, size, bold=True)
            _run(paragraph, text, size, italic=True)
        else:  # para, skill, bullet
            if prefix:
                _run(paragraph, prefix, size, bold=True)
            _run(paragraph, text, size)
            if kind == "bullet":
                fmt.left_indent, fmt.first_line_indent = Inches(0.2), Inches(-0.15)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def convert_to_pdf(docx_path: Path) -> Optional[Path]:
    """Convert .docx -> .pdf. Returns None (and logs why) if no converter is available.

    Windows/macOS: docx2pdf (needs Microsoft Word). Linux, or if docx2pdf fails: LibreOffice headless
    (`soffice --headless --convert-to pdf resume.docx`).
    """
    docx_path = Path(docx_path)
    pdf_path = docx_path.with_suffix(".pdf")

    if sys.platform in ("win32", "darwin"):
        try:
            from docx2pdf import convert

            convert(str(docx_path), str(pdf_path))
            if pdf_path.exists():
                return pdf_path
        except Exception as exc:  # Word not installed, COM error, ...
            logger.warning("docx2pdf failed (%s); trying LibreOffice", exc)

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent),
                            str(docx_path)], check=True, capture_output=True, timeout=180)
            if pdf_path.exists():
                return pdf_path
        except (subprocess.SubprocessError, OSError) as exc:
            logger.warning("LibreOffice conversion failed: %s", exc)
    else:
        logger.warning("No PDF converter found (Word via docx2pdf, or LibreOffice 'soffice').")
    return None
