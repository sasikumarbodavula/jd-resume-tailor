"""Skill checkboxes, editable draft + instant ATS rescoring, resume upload, friendly LLM errors."""
import base64

import pytest
from docx import Document

from app import config, llm, main
from app.guardrail import apply_confirmed_skills, check_tailored_resume, find_confirmed_in_bullets
from app.models.schemas import BaseResume, JDRequirements, TailoredResume
from app.nodes.gap_analysis import gap_analysis_node, selectable_skills
from app.resume_parser import extract_text, parse_resume, verify_parsed_resume
from app.skill_utils import format_skill_lines, parse_skill_lines, split_top_level
from tests.conftest import tailored_from
from tests.test_pipeline_e2e import JD, JD_REQUIREMENTS, good_tailoring, make_fake_llm


@pytest.fixture
def sandbox(monkeypatch, tmp_path):
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path / "out"))
    monkeypatch.setattr(config, "SESSIONS_DIR", tmp_path / "sessions")
    monkeypatch.setattr("app.nodes.export.convert_to_pdf", lambda p: None)
    return tmp_path


# ------------------------------------------------------------ skill helpers
def test_skill_line_parsing_keeps_parentheses_together():
    assert split_top_level("A (x, y), B") == ["A (x, y)", "B"]
    groups = [{"category": "Backend", "items": ["Python", "Authentication (API Keys, OAuth)"]}]
    assert parse_skill_lines(format_skill_lines(groups)) == groups
    assert parse_skill_lines("just, some, skills") == [{"category": "Skills", "items": ["just", "some", "skills"]}]
    assert parse_skill_lines("\n  \n") == []


# ------------------------------------------------------------ checkbox / confirmed skills
def test_selectable_skills_lists_only_what_resume_lacks(base_resume):
    options = selectable_skills(JD_REQUIREMENTS.model_dump(), base_resume["all_known_skills"])
    names = {o["skill"] for o in options}
    assert {"Kubernetes", "Docker", "Snowflake"} <= names
    assert not names & {"Python", "RAG", "LangGraph", "FastAPI"}
    assert len(names) == len(options)  # de-duplicated


def test_apply_confirmed_skills_makes_them_known_and_matched(base_resume):
    effective = apply_confirmed_skills(base_resume, ["Docker", "python", "  "])
    assert effective["confirmed_skills"] == ["Docker"]  # 'python' already owned, blank ignored
    assert effective["skills"][-1] == {"category": "Additional Skills", "items": ["Docker"]}
    gap = gap_analysis_node({"jd_requirements": JD_REQUIREMENTS.model_dump(), "base_resume": effective})["gap_analysis"]
    assert "docker" in gap["matched_skills"] and gap["missing_entirely_skills"] == ["kubernetes"]
    assert base_resume["skills"][-1]["category"] == "Web & Dev Tools"  # original not mutated


def test_confirmed_skill_allowed_in_summary_and_skills_but_not_in_bullets(base_resume):
    effective = apply_confirmed_skills(base_resume, ["Docker"])
    ok = tailored_from(effective, summary="GenAI engineer using Python and Docker.")
    assert check_tailored_resume(effective, ok, [], ["Docker"]) == []
    bad = tailored_from(effective)
    bad["experience_bullets"][0]["text"] = "Containerised the YOLOv5 service with Docker."
    assert find_confirmed_in_bullets(bad, ["Docker"]) == ["docker"]
    assert any("must not appear inside bullets" in p for p in check_tailored_resume(effective, bad, [], ["Docker"]))
    unticked = tailored_from(effective, summary="Uses Kubernetes daily.")
    assert "kubernetes" in check_tailored_resume(effective, unticked, [], ["Docker"])


def test_ticking_skills_raises_ats_and_before_score_stays_original(base_resume, monkeypatch, sandbox):
    calls = {"jd": 0}

    def llm_stub(schema, prompt, system=None):
        if schema is JDRequirements:
            calls["jd"] += 1
            return JD_REQUIREMENTS
        t = tailored_from(base_resume, summary="GenAI engineer with Python, RAG, LangGraph, FastAPI and Docker.")
        t["skills"].append({"category": "Additional Skills", "items": ["Docker", "Kubernetes"]})
        return TailoredResume.model_validate(t)

    monkeypatch.setattr(llm, "invoke_structured", llm_stub)
    analysis = main.analyze(main.AnalyzeRequest(jd_text=JD, resume_id="resume"))
    assert analysis["ats_score_before"] == 75.0
    assert {"Kubernetes", "Docker"} <= {o["skill"] for o in analysis["selectable_skills"]}

    draft = main.tailor(main.TailorRequest(jd_text=JD, resume_id="resume", session_id=analysis["session_id"],
                                           confirmed_skills=["Docker", "Kubernetes"]))
    assert calls["jd"] == 1  # JD parsed once, reused by /tailor
    assert draft["session_id"] == analysis["session_id"] and draft["draft_version"] == 1
    assert draft["ats_score_before"] == 75.0  # still measured on the ORIGINAL resume
    assert draft["ats_score_after"] == 100.0
    assert draft["confirmed_skills"] == ["Docker", "Kubernetes"]


def test_unticked_skill_is_still_blocked_end_to_end(base_resume, monkeypatch, sandbox):
    def llm_stub(schema, prompt, system=None):
        if schema is JDRequirements:
            return JD_REQUIREMENTS
        return TailoredResume.model_validate(tailored_from(base_resume, summary="Experienced with Docker."))

    monkeypatch.setattr(llm, "invoke_structured", llm_stub)
    with pytest.raises(main.HTTPException) as err:
        main.tailor(main.TailorRequest(jd_text=JD, resume_id="resume", confirmed_skills=[]))
    assert err.value.status_code == 422 and "Guardrail" in err.value.detail


# ------------------------------------------------------------ edit + instant ATS update
def test_editing_updates_ats_score_immediately(base_resume, monkeypatch, sandbox):
    monkeypatch.setattr(llm, "invoke_structured", make_fake_llm(base_resume, [good_tailoring(base_resume)]))
    draft = main.tailor(main.TailorRequest(jd_text=JD, resume_id="resume"))
    assert draft["ats_score_after"] == 75.0

    edited = TailoredResume.model_validate(draft["tailored_resume"])
    edited.summary = "GenAI engineer with Python, RAG, LangGraph, FastAPI, Docker and Kubernetes."
    edited.experience_bullets[1].skills_used = ["Kubernetes"]  # tampering with tags is ignored
    result = main.update(main.UpdateRequest(session_id=draft["session_id"], tailored_resume=edited))

    assert result["ats_score_after"] == 100.0 and result["ats_score_before"] == 75.0
    assert {"'docker'", "'kubernetes'"} <= {w.split(" is not")[0] for w in result["warnings"]}
    assert all(k["in_tailored"] for k in result["keywords"] if k["keyword"] in ("Docker", "Kubernetes"))
    saved = main._load_session(draft["session_id"])["tailored_resume"]
    assert "Kubernetes" in saved["summary"]
    assert saved["experience_bullets"][1]["skills_used"] == base_resume["experience_bullets"][1]["skills_used"]


def test_export_uses_edited_text_and_update_invalidates_old_export(base_resume, monkeypatch, sandbox):
    monkeypatch.setattr(llm, "invoke_structured", make_fake_llm(base_resume, [good_tailoring(base_resume)]))
    draft = main.tailor(main.TailorRequest(jd_text=JD, resume_id="resume"))
    sid = draft["session_id"]
    edited = TailoredResume.model_validate(draft["tailored_resume"])
    edited.summary = "My hand-edited summary about Python."
    main.update(main.UpdateRequest(session_id=sid, tailored_resume=edited))
    done = main.approve(main.ApproveRequest(session_id=sid, approved=True))
    text = "\n".join(p.text for p in Document(done["output_paths"]["docx"]).paragraphs)
    assert "My hand-edited summary about Python." in text

    edited.summary = "Second edit."
    main.update(main.UpdateRequest(session_id=sid, tailored_resume=edited))
    assert main._load_session(sid)["output_paths"] is None


def test_update_and_approve_before_tailor_are_rejected(base_resume, monkeypatch, sandbox):
    monkeypatch.setattr(llm, "invoke_structured", lambda s, p, system=None: JD_REQUIREMENTS)
    analysis = main.analyze(main.AnalyzeRequest(jd_text=JD, resume_id="resume"))
    for call in (lambda: main.approve(main.ApproveRequest(session_id=analysis["session_id"], approved=True)),
                 lambda: main.update(main.UpdateRequest(session_id=analysis["session_id"],
                                                        tailored_resume=TailoredResume.model_validate(
                                                            tailored_from(base_resume))))):
        with pytest.raises(main.HTTPException) as err:
            call()
        assert err.value.status_code == 400


# ------------------------------------------------------------ resume upload
SAMPLE = ("JANE DOE\nData Analyst | jane@example.com | Pune, India\nSUMMARY\nAnalyst with 3 years of experience "
          "in reporting and dashboards using SQL and Excel.\nSKILLS\nSQL, Excel, Tableau\nEXPERIENCE\n"
          "Acme Corp - Analyst (2022-2024)\n- Built weekly dashboards in Tableau for 12 stakeholders.\n"
          "EDUCATION\nB.Sc Statistics, Pune University (2021)\n")


def fake_parsed():
    return BaseResume.model_validate({
        "name": "JANE DOE", "title": "Data Analyst", "summary": "Analyst with 3 years of experience in reporting.",
        "skills": [{"category": "Tools", "items": ["SQL", "Excel", "Tableau", "Kubernetes"]}],
        "experience_bullets": [{"id": "x", "text": "Built weekly dashboards in Tableau for 12 stakeholders.",
                                "skills_used": ["Tableau", "Python"]}],
        "projects": [], "education": "B.Sc Statistics, Pune University (2021)", "contact": "jane@example.com"})


def test_extract_text_docx_txt_and_errors(tmp_path):
    path = tmp_path / "r.docx"
    doc = Document()
    for line in SAMPLE.splitlines():
        doc.add_paragraph(line)
    doc.save(path)
    assert "Built weekly dashboards" in extract_text("r.docx", path.read_bytes())
    assert "JANE DOE" in extract_text("r.txt", SAMPLE.encode())
    with pytest.raises(ValueError, match="Unsupported"):
        extract_text("r.exe", b"x" * 500)
    with pytest.raises(ValueError, match="enough text"):
        extract_text("r.txt", b"short")


def test_parser_drops_invented_skills_and_tags():
    cleaned, warnings = verify_parsed_resume(fake_parsed().model_dump(), SAMPLE)
    assert cleaned["skills"][0]["items"] == ["SQL", "Excel", "Tableau"]  # Kubernetes was not in the file
    assert cleaned["experience_bullets"][0]["skills_used"] == ["Tableau"]  # Python was not in the file
    assert cleaned["experience_bullets"][0]["id"] == "exp1_b1"
    assert "Kubernetes" in warnings[0] and "Python" in warnings[0]


def test_upload_parse_then_save_then_use(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(llm, "invoke_structured", lambda s, p, system=None: fake_parsed())
    payload = main.ParseRequest(filename="jane.txt", content_base64=base64.b64encode(SAMPLE.encode()).decode())
    parsed = main.parse_uploaded_resume(payload)
    assert "kubernetes" not in [s.lower() for s in parsed["resume"]["all_known_skills"]]

    saved = main.save_resume(main.SaveResumeRequest(resume_id="jane", resume=parsed["resume"]))
    assert saved["saved"] and (tmp_path / "jane.json").exists()
    assert main.list_resumes() == {"resumes": ["jane"]}
    assert main.get_resume("jane")["name"] == "JANE DOE"
    with pytest.raises(main.HTTPException) as err:
        main.save_resume(main.SaveResumeRequest(resume_id="jane", resume=parsed["resume"]))
    assert err.value.status_code == 409
    main.save_resume(main.SaveResumeRequest(resume_id="jane", resume=parsed["resume"], overwrite=True))
    with pytest.raises(main.HTTPException) as err:
        main.save_resume(main.SaveResumeRequest(resume_id="../evil", resume=parsed["resume"]))
    assert err.value.status_code == 400
    with pytest.raises(main.HTTPException) as err:
        main.save_resume(main.SaveResumeRequest(resume_id="bad", resume={"name": "x"}))
    assert err.value.status_code == 422
    with pytest.raises(main.HTTPException) as err:
        main.parse_uploaded_resume(main.ParseRequest(filename="a.txt", content_base64="!!notbase64!!"))
    assert err.value.status_code == 400


# ------------------------------------------------------------ friendly LLM errors
class _QuotaLLM:
    calls = 0

    def with_structured_output(self, schema):
        return self

    def invoke(self, messages):
        _QuotaLLM.calls += 1
        raise RuntimeError("429 You exceeded your current quota ... gemini-3.8-flash")


def test_quota_error_is_friendly_and_not_retried(monkeypatch):
    monkeypatch.setattr(llm, "get_llm", lambda: _QuotaLLM())
    with pytest.raises(llm.LLMError, match="rate limit"):
        llm.invoke_structured(JDRequirements, "x")
    assert _QuotaLLM.calls == 1
    with pytest.raises(main.HTTPException) as err:
        main._run(lambda: llm.invoke_structured(JDRequirements, "x"))
    assert err.value.status_code == 502


def test_ticked_skill_is_never_lost_by_the_llm(base_resume):
    from app.guardrail import enforce_ground_truth
    t = tailored_from(base_resume)  # LLM output without the ticked skill
    fixed = enforce_ground_truth(base_resume, t, ensure_skills=["Docker", "python"]).model_dump()
    assert fixed["skills"][-1] == {"category": "Additional Skills", "items": ["Docker"]}
